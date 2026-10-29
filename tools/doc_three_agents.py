from docx import Document
import re

path = r"D:\Program Files (x86)\xwechat_files\wxid_qt0mzs5wekag22_65ad\msg\file\2026-10\EAG-Guard阶段性研究报告.docx"
doc = Document(path)

keys = ["三个", "三票", "投票", "表决", "多数", "审核表", "裁判表", "评审表", "多智能体",
        "多个裁判", "两名", "三位", "ensemble", "投票"]
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    if any(k in t for k in keys):
        print(f"[{i}] {t[:400]}\n")

print("=== 表格里是否有相关内容 ===")
for ti, table in enumerate(doc.tables):
    for row in table.rows:
        cells = [c.text.strip() for c in row.cells]
        line = " | ".join(cells)
        if any(k in line for k in keys + ["裁判", "审核"]):
            print(f"表{ti}: {line[:200]}")
