# -*- coding: utf-8 -*-
"""Print a structural overview of a .docx report."""
import sys
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

SRC = sys.argv[1]


def iter_block_items(parent):
    from docx.oxml.ns import qn
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)


doc = Document(SRC)
print("SECTIONS:", len(doc.sections))
for i, s in enumerate(doc.sections):
    print(f"  sec{i}: page={s.page_width}x{s.page_height} margins L{s.left_margin} R{s.right_margin}")
print("STYLES USED:")
counts = {}
for p in doc.paragraphs:
    counts[p.style.name] = counts.get(p.style.name, 0) + 1
for k, v in sorted(counts.items(), key=lambda x: -x[1]):
    print(f"  {k}: {v}")
print("TABLES:", len(doc.tables))
print("INLINE SHAPES:", len(doc.inline_shapes))
chars = sum(len(p.text) for p in doc.paragraphs)
print("PARAGRAPH CHARS:", chars)
print("---- OUTLINE ----")
for i, block in enumerate(iter_block_items(doc)):
    if isinstance(block, Paragraph):
        t = block.text.strip()
        if t and (block.style.name.startswith('Heading') or block.style.name in ('Title', 'Subtitle')):
            print(f"[{i}] <{block.style.name}> {t[:100]}")
    else:
        print(f"[{i}] <TABLE {len(block.rows)}x{len(block.columns)}>")
